import struct
from pathlib import Path

import pytest
from PIL import Image
from pypdf import PdfWriter

from app.ingestion.cleaner import clean_text
from app.ingestion.loader import (
    MAIN_MASTER,
    PERSIST_DIRECTORY,
    SLIDE,
    SLIDE_LIST,
    SLIDE_PERSIST,
    TEXT_BYTES,
    TEXT_CHARS,
    clean_title,
    load_document,
    parse_ppt_stream,
)


def test_load_text_file(tmp_path):
    file = tmp_path / "notes.txt"
    file.write_text("hello world", encoding="utf-8")

    doc = load_document(file)

    assert doc.text == "hello world"
    assert doc.metadata["source"] == "notes.txt"


def test_unsupported_extension(tmp_path):
    file = tmp_path / "video.mp4"
    file.write_bytes(b"")
    with pytest.raises(ValueError):
        load_document(file)


def test_clean_text():
    raw = "Retrieval   aug-\nmented\n\n\n\ngeneration "
    assert clean_text(raw) == "Retrieval augmented\n\ngeneration"


def test_load_pptx_keeps_slide_numbers(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    for title in ("Hydatidose hépatique", "Traitement chirurgical"):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"Contenu: {title}"
    deck.save(tmp_path / "cours.pptx")

    doc = load_document(tmp_path / "cours.pptx")

    assert [p.number for p in doc.pages] == [1, 2]
    assert "Traitement chirurgical" in doc.pages[1].text


def test_load_docx_with_table(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("Score de Glasgow")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Ouverture des yeux", "4"
    document.save(tmp_path / "fiche.docx")

    text = load_document(tmp_path / "fiche.docx").text

    assert "Score de Glasgow" in text
    assert "Ouverture des yeux | 4" in text


def _record(rec_type: int, body: bytes, container: bool = False, instance: int = 0) -> bytes:
    ver_inst = (instance << 4) | (0xF if container else 0)
    return struct.pack("<HHI", ver_inst, rec_type, len(body)) + body


def test_parse_legacy_ppt_stream():
    slide_list = _record(
        SLIDE_LIST,
        _record(SLIDE_PERSIST, b"\0" * 20)
        + _record(TEXT_CHARS, "Maladie de Horton\rArtérite".encode("utf-16-le"))
        + _record(SLIDE_PERSIST, b"\0" * 20)
        + _record(TEXT_BYTES, "Biopsie temporale".encode("latin-1")),
        container=True,
    )
    master = _record(MAIN_MASTER, _record(TEXT_CHARS, "Cliquez".encode("utf-16-le")), True)
    stream = _record(0x03E8, slide_list + master, container=True)  # DocumentContainer

    pages = parse_ppt_stream(stream)

    assert [(p.number, p.text) for p in pages] == [
        (1, "Maladie de Horton\nArtérite"),
        (2, "Biopsie temporale"),
    ]


def test_parse_legacy_ppt_text_inside_slide_records():
    # Presentation order: persist id 5 is slide 1, persist id 3 is slide 2.
    slide_list = _record(
        SLIDE_LIST,
        _record(SLIDE_PERSIST, struct.pack("<I", 5) + b"\0" * 16)
        + _record(SLIDE_PERSIST, struct.pack("<I", 3) + b"\0" * 16),
        container=True,
    )
    document = _record(0x03E8, slide_list, container=True)
    second = _record(SLIDE, _record(TEXT_CHARS, "Deuxième".encode("utf-16-le")), True)
    first = _record(SLIDE, _record(TEXT_CHARS, "Première".encode("utf-16-le")), True)
    directory = struct.pack("<II", 1 << 20 | 3, len(document)) + struct.pack(
        "<II", 1 << 20 | 5, len(document) + len(second)
    )
    stream = document + second + first + _record(PERSIST_DIRECTORY, directory)

    pages = parse_ppt_stream(stream)

    assert [(p.number, p.text) for p in pages] == [(1, "Première"), (2, "Deuxième")]


def test_clean_title():
    assert clean_title(Path("1-[25-26]Physiologie du vieillissement .pdf")) == (
        "Physiologie du vieillissement"
    )
    assert clean_title(Path("6- TRIAGE 2025.pptx.pdf")) == "TRIAGE 2025"


def fake_ocr(image: Image.Image) -> str:
    return f"texte lu par OCR ({image.width}x{image.height})"


def test_scanned_pdf_page_goes_through_ocr(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)  # no text layer, like a scan
    with open(tmp_path / "scan.pdf", "wb") as f:
        writer.write(f)

    without_ocr = load_document(tmp_path / "scan.pdf")
    with_ocr = load_document(tmp_path / "scan.pdf", ocr=fake_ocr)

    assert without_ocr.pages == []
    assert with_ocr.pages[0].number == 1
    assert with_ocr.pages[0].text.startswith("texte lu par OCR")


def test_image_only_slide_goes_through_ocr(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    Image.new("RGB", (800, 600), "white").save(tmp_path / "scan.png")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])  # blank layout
    slide.shapes.add_picture(str(tmp_path / "scan.png"), 0, 0, width=Inches(8))
    deck.save(tmp_path / "scans.pptx")

    doc = load_document(tmp_path / "scans.pptx", ocr=fake_ocr)

    assert "texte lu par OCR (800x600)" in doc.pages[0].text


def test_photo_is_read_with_ocr(tmp_path):
    Image.new("RGB", (640, 480), "white").save(tmp_path / "notes.jpg")

    assert load_document(tmp_path / "notes.jpg").pages == []
    assert load_document(tmp_path / "notes.jpg", ocr=fake_ocr).text.endswith("(640x480)")
